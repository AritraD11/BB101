"""Build the printable A4 attendance sheet for BB 101 Tutorial Batch T1.

Reads the T1 roster out of the course list workbook and writes a one-page
A4 sheet with seven blank date columns for marking P / A by hand.
"""

import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

SOURCE = sys.argv[1] if len(sys.argv) > 1 else "CourseList_Tutorail.xlsx"
OUTPUT = sys.argv[2] if len(sys.argv) > 2 else "BB101_T1_Attendance.xlsx"

N_DATE_COLS = 7

# ---------------------------------------------------------------- read roster
src = openpyxl.load_workbook(SOURCE)["T1"]
roster = []
for row in src.iter_rows(min_row=3, max_row=src.max_row, values_only=True):
    roll, name = row[1], row[2]
    if roll is None or name is None:
        continue
    roster.append((str(roll).strip(), str(name).strip()))

# --------------------------------------------------------------- build sheet
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "T1 Attendance"

FONT = "Arial"
thin = Side(style="thin", color="808080")
medium = Side(style="medium", color="000000")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
header_fill = PatternFill("solid", fgColor="D9D9D9")
date_fill = PatternFill("solid", fgColor="F2F2F2")

last_col = 3 + N_DATE_COLS  # Sr | Roll | Name | 7 dates
last_letter = get_column_letter(last_col)

# Title block
ws.merge_cells(f"A1:{last_letter}1")
ws["A1"] = "BB 101  —  BIOLOGY   |   ATTENDANCE SHEET"
ws["A1"].font = Font(name=FONT, size=15, bold=True)
ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

ws.merge_cells(f"A2:{last_letter}2")
ws["A2"] = "Tutorial Batch T1        Room: LT 206        TA: Aritra"
ws["A2"].font = Font(name=FONT, size=11, bold=True)
ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

# Header rows 3-4: Sr/Roll/Name span both, date band on top of blank date cells
HDR_TOP, HDR_BOT = 3, 4
for col, label in ((1, "Sr."), (2, "Roll No."), (3, "Name of Student")):
    ws.merge_cells(start_row=HDR_TOP, start_column=col, end_row=HDR_BOT, end_column=col)
    c = ws.cell(row=HDR_TOP, column=col, value=label)
    c.font = Font(name=FONT, size=10, bold=True)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

# Left unmerged (only "centerContinuous" across the date cells) so the
# vertical dividers between date columns still show through the header,
# matching the printed PDF.
for col in range(4, last_col + 1):
    c = ws.cell(row=HDR_TOP, column=col, value="DATE" if col == 4 else None)
    c.font = Font(name=FONT, size=10, bold=True)
    c.alignment = Alignment(horizontal="centerContinuous", vertical="center")

for col in range(4, last_col + 1):
    c = ws.cell(row=HDR_BOT, column=col, value=None)
    c.fill = date_fill

for row in (HDR_TOP, HDR_BOT):
    for col in range(1, last_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.border = box
        if cell.fill.fgColor.rgb != "00F2F2F2":
            cell.fill = header_fill

# Student rows
first_data = HDR_BOT + 1
for i, (roll, name) in enumerate(roster):
    r = first_data + i
    ws.cell(row=r, column=1, value=i + 1).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=r, column=2, value=roll).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=r, column=3, value=name).alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for col in range(1, last_col + 1):
        cell = ws.cell(row=r, column=col)
        cell.border = box
        cell.font = Font(name=FONT, size=10)
    ws.row_dimensions[r].height = 19

last_data = first_data + len(roster) - 1

# Totals row — counts P / p in each date column
total_row = last_data + 1
ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=3)
t = ws.cell(row=total_row, column=1, value="TOTAL PRESENT")
t.font = Font(name=FONT, size=10, bold=True)
t.alignment = Alignment(horizontal="right", vertical="center", indent=1)
for col in range(4, last_col + 1):
    L = get_column_letter(col)
    rng = f"{L}{first_data}:{L}{last_data}"
    # blank until the column is actually marked, so a freshly printed sheet is clean
    c = ws.cell(row=total_row, column=col, value=f'=IF(COUNTA({rng})=0,"",COUNTIF({rng},"P"))')
    c.font = Font(name=FONT, size=10, bold=True)
    c.alignment = Alignment(horizontal="center", vertical="center")
for col in range(1, last_col + 1):
    ws.cell(row=total_row, column=col).border = box
    ws.cell(row=total_row, column=col).fill = header_fill
ws.row_dimensions[total_row].height = 20

# Footer
sig_row = total_row + 2
ws.merge_cells(start_row=sig_row, start_column=1, end_row=sig_row, end_column=last_col)
s = ws.cell(row=sig_row, column=1, value=f"Total students on roll: {len(roster)}")
s.font = Font(name=FONT, size=9, color="404040")
s.alignment = Alignment(horizontal="left", vertical="center")

# Outer border around the whole grid
for col in range(1, last_col + 1):
    top = ws.cell(row=HDR_TOP, column=col)
    top.border = Border(left=top.border.left, right=top.border.right, top=medium, bottom=top.border.bottom)
    bot = ws.cell(row=total_row, column=col)
    bot.border = Border(left=bot.border.left, right=bot.border.right, top=bot.border.top, bottom=medium)
for row in range(HDR_TOP, total_row + 1):
    lft = ws.cell(row=row, column=1)
    lft.border = Border(left=medium, right=lft.border.right, top=lft.border.top, bottom=lft.border.bottom)
    rgt = ws.cell(row=row, column=last_col)
    rgt.border = Border(left=rgt.border.left, right=medium, top=rgt.border.top, bottom=rgt.border.bottom)
# separator between Name and the date block
for row in range(HDR_TOP, total_row + 1):
    c = ws.cell(row=row, column=4)
    c.border = Border(left=medium, right=c.border.right, top=c.border.top, bottom=c.border.bottom)

# ------------------------------------------------------------ column widths
ws.column_dimensions["A"].width = 4.5
ws.column_dimensions["B"].width = 11.5
ws.column_dimensions["C"].width = 32
for col in range(4, last_col + 1):
    ws.column_dimensions[get_column_letter(col)].width = 6.6

ws.row_dimensions[1].height = 22
ws.row_dimensions[2].height = 18
ws.row_dimensions[HDR_TOP].height = 17
ws.row_dimensions[HDR_BOT].height = 30

# --------------------------------------------------------------- print setup
ws.page_setup.orientation = "portrait"
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
ws.page_margins.left = 0.4
ws.page_margins.right = 0.4
ws.page_margins.top = 0.4
ws.page_margins.bottom = 0.4
ws.page_margins.header = 0.2
ws.page_margins.footer = 0.2
ws.print_options.horizontalCentered = True
ws.print_title_rows = f"{HDR_TOP}:{HDR_BOT}"
ws.freeze_panes = ws.cell(row=first_data, column=3)

wb.save(OUTPUT)
print(f"wrote {OUTPUT}: {len(roster)} students, {N_DATE_COLS} date columns, rows {first_data}-{last_data}")
