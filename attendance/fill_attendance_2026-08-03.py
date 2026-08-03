"""Record the 03.08.2026 attendance (from the marked printout photo) into the
first date column of the T1 attendance workbook.

Blank in the source list below = no mark on the physical sheet (left as-is,
per TA's call, rather than inferred as absent).
"""

import openpyxl
from openpyxl.styles import Alignment, Font

FILE = "BB101_T1_Attendance.xlsx"
FONT = "Arial"
DATE_COL = 4       # column D = first date column
DATE_ROW = 4       # HDR_BOT row, where a date is handwritten under "DATE"
DATE_LABEL = "03.08.2026"
FIRST_DATA_ROW = 5  # Sr. 1

# Sr. 1-31, in roster order, exactly as marked on the printed sheet.
MARKS = [
    None, None, "P", "P", "P", "P", "P", None, "P", "P",
    "P", "P", "P", "P", "P", "P", "P", "P", "P", "P",
    "P", "P", "P", "P", "P", "P", "P", "P", "P", "P", "P",
]

wb = openpyxl.load_workbook(FILE)
ws = wb.active

date_cell = ws.cell(row=DATE_ROW, column=DATE_COL, value=DATE_LABEL)
date_cell.font = Font(name=FONT, size=9, bold=True)
date_cell.alignment = Alignment(horizontal="center", vertical="center")

for i, mark in enumerate(MARKS):
    if mark is None:
        continue
    cell = ws.cell(row=FIRST_DATA_ROW + i, column=DATE_COL, value=mark)
    cell.alignment = Alignment(horizontal="center", vertical="center")

wb.save(FILE)
present = sum(1 for m in MARKS if m == "P")
blank = sum(1 for m in MARKS if m is None)
print(f"wrote {DATE_LABEL} into column D: {present} present, {blank} unmarked, {len(MARKS)} total")
