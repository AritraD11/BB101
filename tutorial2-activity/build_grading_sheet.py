"""Build the printable A4 grading sheet for the Tutorial 2 group activity.

One page, T1 roster (31 students), with a Group column and three graded
criteria (Understanding / Clarity / Insight, each out of 5) plus a Total.
"""

import sys

import openpyxl
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas

SOURCE = sys.argv[1] if len(sys.argv) > 1 else "../attendance/CourseList_Tutorail.xlsx"
OUTPUT = sys.argv[2] if len(sys.argv) > 2 else "BB101_T1_GroupActivity_Grading.pdf"

src = openpyxl.load_workbook(SOURCE)["T1"]
roster = []
for row in src.iter_rows(min_row=3, max_row=src.max_row, values_only=True):
    roll, name = row[1], row[2]
    if roll is None or name is None:
        continue
    roster.append((str(roll).strip(), str(name).strip()))

PAGE_W, PAGE_H = A4
M_L = M_R = 10 * mm
M_T = 11 * mm
M_B = 10 * mm
GRID_W = PAGE_W - M_L - M_R

W_SR = 9 * mm
W_ROLL = 22 * mm
W_GROUP = 16 * mm
W_CRIT = 22 * mm   # Understanding / Clarity / Insight
W_TOTAL = 20 * mm
W_NAME = GRID_W - W_SR - W_ROLL - W_GROUP - 3 * W_CRIT - W_TOTAL

COL_X = [M_L]
for w in (W_SR, W_ROLL, W_NAME, W_GROUP, W_CRIT, W_CRIT, W_CRIT, W_TOTAL):
    COL_X.append(COL_X[-1] + w)
# COL_X: [0]left, [1]Sr/Roll, [2]Roll/Name, [3]Name/Group, [4]Group/Und,
#        [5]Und/Clar, [6]Clar/Ins, [7]Ins/Total, [8]right edge

ROW_H = 7.2 * mm
HDR_H = 14 * mm

INK = colors.black
GRID = colors.HexColor("#8C8C8C")
SHADE = colors.HexColor("#DCDCDC")

c = canvas.Canvas(OUTPUT, pagesize=A4)
c.setTitle("BB 101 - Tutorial 2 Group Activity Grading Sheet")

y = PAGE_H - M_T
c.setFont("Helvetica-Bold", 14)
c.drawCentredString(PAGE_W / 2, y, "BB 101  —  BIOLOGY   |   GROUP ACTIVITY GRADING SHEET")
y -= 6.0 * mm
c.setFont("Helvetica-Bold", 10.5)
c.drawCentredString(PAGE_W / 2, y, "Tutorial Batch T1        Room: LT 206        TA: Aritra")
y -= 4.6 * mm
c.setFont("Helvetica-Oblique", 8)
c.setFillColor(colors.HexColor("#555555"))
c.drawCentredString(PAGE_W / 2, y, "Each criterion out of 5  ·  Understanding / Clarity / Insight  ·  Total out of 15")
c.setFillColor(INK)
y -= 6.5 * mm

grid_top = y
hdr_bot = grid_top - HDR_H

c.setFillColor(SHADE)
c.rect(M_L, hdr_bot, GRID_W, HDR_H, stroke=0, fill=1)
c.setFillColor(INK)

headers = ["Sr.", "Roll No.", "Name of Student", "Group", "Underst.\n/5", "Clarity\n/5", "Insight\n/5", "Total\n/15"]
c.setFont("Helvetica-Bold", 8.3)
for i, h in enumerate(headers):
    cx = (COL_X[i] + COL_X[i+1]) / 2
    lines = h.split("\n")
    ty = hdr_bot + HDR_H/2 + (3.0 * mm if len(lines) > 1 else 1.6 * mm)
    for ln in lines:
        c.drawCentredString(cx, ty, ln)
        ty -= 3.6 * mm

row_top = hdr_bot
c.setFont("Helvetica", 9)
for i, (roll, name) in enumerate(roster):
    ry = row_top - i * ROW_H
    base = ry - ROW_H + 2.3 * mm
    if i % 2 == 1:
        c.setFillColor(colors.HexColor("#FAFAFA"))
        c.rect(M_L, ry - ROW_H, GRID_W, ROW_H, stroke=0, fill=1)
        c.setFillColor(INK)
    c.setFont("Helvetica", 9)
    c.drawCentredString((COL_X[0]+COL_X[1])/2, base, str(i + 1))
    c.drawCentredString((COL_X[1]+COL_X[2])/2, base, roll)
    label = name
    while c.stringWidth(label, "Helvetica", 9) > W_NAME - 5 * mm and len(label) > 4:
        label = label[:-2]
    c.drawString(COL_X[2] + 2.5 * mm, base, label)

grid_bot = row_top - len(roster) * ROW_H

c.setStrokeColor(GRID)
c.setLineWidth(0.4)
for i in range(len(roster) + 1):
    yy = row_top - i * ROW_H
    c.line(M_L, yy, PAGE_W - M_R, yy)
for x in COL_X[1:-1]:
    c.line(x, grid_bot, x, grid_top)

c.setStrokeColor(INK)
c.setLineWidth(1.1)
c.rect(M_L, grid_bot, GRID_W, grid_top - grid_bot, stroke=1, fill=0)
c.line(COL_X[3], grid_bot, COL_X[3], grid_top)
c.line(M_L, hdr_bot, PAGE_W - M_R, hdr_bot)

fy = grid_bot - 8 * mm
c.setFont("Helvetica", 8.5)
c.setFillColor(colors.HexColor("#404040"))
c.drawString(M_L, fy, f"Total students on roll: {len(roster)}")
c.setFillColor(INK)

ky = fy - 8 * mm
c.setFont("Helvetica-Oblique", 8)
c.setFillColor(colors.HexColor("#555555"))
c.drawString(M_L, ky, "Understanding — got the core idea right   ·   Clarity — followable in under 2 min   ·   Insight — a genuine reaction, not a summary")
c.setFillColor(INK)

c.showPage()
c.save()
print(f"wrote {OUTPUT}: {len(roster)} students, 1 A4 page")
